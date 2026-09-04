#!/usr/bin/env python3
"""Fills a workbook with a plausible year-in-progress so the listing images
show a living tracker rather than an empty grid."""
import datetime as dt
import os
import random
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(os.path.dirname(HERE), "dist")
START = dt.date(2026, 1, 5)

# per habit: (probability of a tick, probability an explicit ✗ is used)
# deliberately imperfect: a tracker that shows a flawless year teaches the
# buyer nothing about what the heat map looks like when they slip
PROFILE = [(0.72, 0.35), (0.60, 0.25), (0.55, 0.30), (0.80, 0.15),
           (0.50, 0.40), (0.66, 0.30), (0.58, 0.25)]

# two extra habits so the listing screenshots show a populated grid rather
# than five rows and a void; the shipped file still ships with five
EXTRA_HABITS = [("Meditate 10 min", 4), ("Inbox to zero", 5)]


def fill(src, dst, seed=7):
    rnd = random.Random(seed)
    today_idx = (dt.date.today() - START).days + 1
    wb = openpyxl.load_workbook(src)
    setup = wb["Start Here"]
    for i, (name, goal) in enumerate(EXTRA_HABITS):
        setup.cell(row=11 + 5 + i, column=3, value=name)
        setup.cell(row=11 + 5 + i, column=5, value=goal)
    sheets = {}
    for h, (p_do, p_miss) in enumerate(PROFILE, start=1):
        run = 0
        for i in range(1, min(today_idx, 364) + 1):
            # a habit in a streak is likelier to keep going
            p = min(0.93, p_do + 0.04 * min(run, 3))
            hit = rnd.random() < p
            # keep habit 4 unbroken over the last fortnight: the hero streak
            if h == 4 and i > today_idx - 14:
                hit = True
            run = run + 1 if hit else 0
            if not hit and rnd.random() > p_miss:
                continue                      # left blank rather than ✗
            w, d = (i - 1) // 7 + 1, (i - 1) % 7
            ws = sheets.setdefault(w, wb["W%02d" % w])
            ws.cell(row=8 + h, column=3 + d, value="✓" if hit else "✗")

    # habit 2 is deliberately left unticked today while its streak is live,
    # so the Today page shows the AT RISK state the listing talks about
    for i in range(today_idx - 5, today_idx):
        w, d = (i - 1) // 7 + 1, (i - 1) % 7
        sheets.setdefault(w, wb["W%02d" % w]).cell(row=10, column=3 + d,
                                                   value="✓")
    w, d = (today_idx - 1) // 7 + 1, (today_idx - 1) % 7
    # ws.cell(value=None) is a no-op in openpyxl, so clear it by assignment
    sheets.setdefault(w, wb["W%02d" % w]).cell(row=10,
                                               column=3 + d).value = None

    cur_week = (min(today_idx, 364) - 1) // 7 + 1
    wb["W%02d" % cur_week]["C24"] = ("Protect the morning block — no phone "
                                     "before the first habit is done.")
    wb["W%02d" % cur_week]["C26"] = ("Two misses both landed on days I "
                                     "skipped breakfast. Not a coincidence.")
    wb.save(dst)
    return dst, cur_week


if __name__ == "__main__":
    for theme in ("Light", "Dark"):
        src = os.path.join(DIST, "Momentum-Habit-System-%s.xlsx" % theme)
        out = os.path.join(sys.argv[1] if len(sys.argv) > 1 else "/tmp",
                           "demo-%s.xlsx" % theme)
        _, cw = fill(src, out)
        print("%-14s -> %s  (current week W%02d)" % (theme, out, cw))
