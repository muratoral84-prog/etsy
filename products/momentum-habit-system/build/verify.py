#!/usr/bin/env python3
"""Fills a generated workbook with a known tick pattern, has LibreOffice do
the real recalculation, and asserts the numbers that come back."""
import datetime as dt
import os
import subprocess
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build as B          # geometry constants come from the generator,
                           # so a layout change can never silently skip a test

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(os.path.dirname(HERE), "dist")
TMP = os.environ.get("VERIFY_TMP", "/tmp/momentum-verify")
START = dt.date(2026, 1, 5)
ERRS = ("#REF!", "#VALUE!", "#NAME?", "#N/A", "#DIV/0!", "#NULL!", "#NUM!",
        "#ERROR!", "Err:")


def recalc(src, outdir):
    os.makedirs(outdir, exist_ok=True)
    dst = os.path.join(outdir, os.path.basename(src))
    if os.path.exists(dst):
        os.remove(dst)
    subprocess.run(
        ["soffice", "--headless", "--norestore",
         "-env:UserInstallation=file:///tmp/lo-profile",
         "--convert-to", "xlsx", "--outdir", outdir, src],
        check=True, capture_output=True, timeout=600)
    return dst


def streaks(days, today_idx):
    at, cur, best = {}, 0, 0
    for i in range(1, 365):
        cur = cur + 1 if i in days else 0
        at[i] = cur
        best = max(best, cur)
    return dict(total=len(days), best=best,
                current=max(at.get(today_idx, 0), at.get(today_idx - 1, 0)))


def check(path):
    today_idx = (dt.date.today() - START).days + 1
    wb = openpyxl.load_workbook(path)
    sh = {}

    def mark(h, days, sym="✓"):
        for i in days:
            w, d = (i - 1) // 7 + 1, (i - 1) % 7
            sh.setdefault(w, wb["W%02d" % w]).cell(row=8 + h, column=3 + d,
                                                   value=sym)

    plan = {
        1: list(range(1, 11)) + list(range(12, 46)),      # gap on day 11
        2: list(range(1, 8)) + list(range(9, 16)),        # explicit miss d8
        3: list(range(1, 101)),                           # 100 straight
        4: list(range(today_idx - 5, today_idx + 1)),     # live streak of 6
        5: list(range(today_idx - 8, today_idx - 5)),     # streak that died
    }
    for h, days in plan.items():
        mark(h, days)
    mark(2, [8], "✗")

    tin = os.path.join(TMP, "in-" + os.path.basename(path))
    os.makedirs(TMP, exist_ok=True)
    wb.save(tin)
    out = recalc(tin, os.path.join(TMP, "out"))

    r = openpyxl.load_workbook(out, data_only=True)
    fails = []

    n_err = 0
    for ws in r.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and any(
                        c.value.startswith(e) for e in ERRS):
                    n_err += 1
                    if n_err < 4:
                        fails.append("formula error %s!%s = %s"
                                     % (ws.title, c.coordinate, c.value))
    if n_err:
        fails.append("%d formula errors total" % n_err)

    d = r["Dashboard"]
    exp_all = {h: streaks(set(days), today_idx) for h, days in plan.items()}
    for h, e in exp_all.items():
        row = 12 + h
        got = dict(total=d["I%d" % row].value, current=d["K%d" % row].value,
                   best=d["M%d" % row].value)
        for k in e:
            if got[k] != e[k]:
                fails.append("habit %d %s: got %r want %r"
                             % (h, k, got[k], e[k]))

    exp_badges = ["✔" if exp_all[3]["best"] >= t else "·"
                  for t in (7, 30, 100, 365)]
    got_badges = [d["%s15" % c].value for c in "OPQR"]
    if got_badges != exp_badges:
        fails.append("badges: got %s want %s" % (got_badges, exp_badges))

    cards = [
        ("C7", sum(v["total"] for v in exp_all.values())),
        ("G7", max(v["best"] for v in exp_all.values())),
        ("K7", sum(1 for v in exp_all.values() if v["current"] > 0)),
    ]
    for coord, want in cards:
        if d[coord].value != want:
            fails.append("card %s: got %r want %r" % (coord, d[coord].value,
                                                      want))

    cur_w = (today_idx - 1) // 7 + 1
    cw = r["W%02d" % cur_w]
    for h in range(1, 6):
        got = cw["N%d" % (B.W_HAB_R0 + h - 1)].value
        if got != exp_all[h]["current"]:
            fails.append("W%02d streak habit %d: got %r want %r (must match "
                         "Dashboard)" % (cur_w, h, got, exp_all[h]["current"]))

    w1 = r["W01"]
    for coord, want in (("K9", 7), ("M9", "✓"), ("N9", 7), ("C20", 3)):
        if w1[coord].value != want:
            fails.append("W01!%s: got %r want %r" % (coord, w1[coord].value,
                                                     want))
    if round(w1["C22"].value, 4) != 0.75:
        fails.append("W01 week score: got %r want 0.75" % w1["C22"].value)

    w52 = r["W52"]                       # a week that has not begun yet
    for coord in ("N9", "C22", "H22"):
        if w52[coord].value not in (None, ""):
            fails.append("W52!%s should be blank, got %r"
                         % (coord, w52[coord].value))
    # first week whose Monday is still in the future -> must be blank
    from openpyxl.utils import get_column_letter as CL
    def map_cell(w, off):
        return "%s%d" % (CL(B.D_C0 + (w - 1) % 13),
                         B.D_IDX_R0 + ((w - 1) // 13) * 4 + off)
    fut = next((w for w in range(1, 53)
                if START + dt.timedelta(days=(w - 1) * 7) > dt.date.today()),
               None)
    if fut:
        coord = map_cell(fut, 1)
        if d[coord].value not in (None, ""):
            fails.append("year map W%02d (%s) should be blank, got %r"
                         % (fut, coord, d[coord].value))
    if round(d[map_cell(1, 1)].value, 4) != 0.75:
        fails.append("year map W01: got %r want 0.75" % d[map_cell(1, 1)].value)

    # no score may ever read above 100%
    for w in range(1, 53):
        v = d[map_cell(w, 1)].value
        if isinstance(v, float) and v > 1.0001:
            fails.append("year map W%02d over 100%%: %r" % (w, v))
    for w in (1, 2, 3):
        v = r["W%02d" % w]["C%d" % B.W_SCORE_R].value
        if isinstance(v, float) and v > 1.0001:
            fails.append("W%02d week score over 100%%: %r" % (w, v))

    # Today page must agree with the Dashboard, not hold a second opinion
    tod = r[B.TODAY]
    today_row = None
    for h in range(1, 6):
        tr = 12 + h - 1
        if tod["D%d" % tr].value != exp_all[h]["current"]:
            fails.append("Today habit %d streak: got %r want %r"
                         % (h, tod["D%d" % tr].value, exp_all[h]["current"]))
        want_status = ("logged" if h == 4 else
                       ("AT RISK" if exp_all[h]["current"] > 0 else "not yet"))
        if tod["E%d" % tr].value != want_status:
            fails.append("Today habit %d status: got %r want %r"
                         % (h, tod["E%d" % tr].value, want_status))
    # habit 3 best=100 -> next unearned badge is 365
    if not str(tod["F14"].value).startswith("365-day badge"):
        fails.append("Today next badge for habit 3: got %r" % tod["F14"].value)

    return fails


def main():
    bad = 0
    for name in sorted(os.listdir(DIST)):
        if not name.endswith(".xlsx"):
            continue
        fails = check(os.path.join(DIST, name))
        print("%-42s %s" % (name, "PASS" if not fails else "FAIL"))
        for f in fails:
            print("      - %s" % f)
        bad += len(fails)
    print("\n%s" % ("all checks passed" if not bad else "%d failures" % bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
